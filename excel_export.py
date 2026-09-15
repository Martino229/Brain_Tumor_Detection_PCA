import os
import pandas as pd
from openpyxl.drawing.image import Image
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, Alignment, numbers

def excel_export(data, output_filename="Report_Diagnostico_Malattia.xlsx", image_paths=None):
    # 1. Creiamo un nuovo workbook
    wb = Workbook()

    # --- FOGLIO 1: DATI NUMERICI ---
    ws1 = wb.active
    ws1.title = "Analisi Statistica"

    # Gestione dinamica dei dati (ora ottimizzata per liste di dizionari)
    if isinstance(data, dict):
        df = pd.DataFrame([data])
    else:
        df = pd.DataFrame(data)

    # Scrittura del DataFrame
    for r in dataframe_to_rows(df, index=False, header=True):
        ws1.append(r)

    # Formattazione estetica dell'intestazione
    for cell in ws1[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Metriche che devono essere formattate visivamente come percentuali
    metri_percentuali = [
        "Varianza Spiegata PCA", "Accuratezza", "Sensibilità (Recall)",
        "Specificità", "Precisione", "F1-Score",
        "Veri Positivi (TP)", "Veri Negativi (TN)",
        "Falsi Positivi (FP)", "Falsi Negativi (FN) [GRAVE]"
    ]

    # Formattazione dinamica per RIGA invece che per colonna
    for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row, min_col=1, max_col=ws1.max_column):
        nome_metrica = row[0].value
        for cell in row[1:]: # Salta la prima colonna (il nome)
            if isinstance(cell.value, (int, float)):
                if nome_metrica in metri_percentuali:
                    cell.number_format = numbers.FORMAT_PERCENTAGE_00
                else:
                    # Per deviazioni standard o numeri interi puri
                    cell.number_format = numbers.FORMAT_NUMBER_00 if isinstance(cell.value, float) else numbers.FORMAT_NUMBER
            cell.alignment = Alignment(horizontal="center")
        row[0].alignment = Alignment(horizontal="left")

    # Auto-adattamento comodo della larghezza colonne
    for col in ws1.columns:
        max_length = max((len(str(cell.value)) for cell in col if cell.value), default=0)
        ws1.column_dimensions[col[0].column_letter].width = max_length + 5

    # --- FOGLIO 2: REPORT GRAFICO ---
    if image_paths and len(image_paths) > 0:
        ws2 = wb.create_sheet(title="Cruscotto Geometrico-Clinico")

        # Griglia ampiamente distanziata per evitare sovrapposizioni (da B a P, riga 3 e 40)
        anchor_cells = ['B3', 'P3', 'B40', 'P40']

        titles = [
            "1. Varianza Cumulativa (Spettro Autovalori)",
            "2. Proiezione Ortogonale nello Spazio PCA (PC1 vs PC2)",
            "3. Analisi Statistica Distanze Minime",
            "4. Matrice di Confusione"
        ]

        for idx, img_path in enumerate(image_paths):
            if os.path.exists(img_path) and idx < 4:
                # Titolo formattato sopra l'immagine
                title_cell = ws2[anchor_cells[idx]].offset(row=-1, column=0)
                title_cell.value = titles[idx]
                title_cell.font = Font(bold=True, size=12)

                # Compressione al 40% per incastrare i grafici nella griglia ed evitare collage
                img = Image(img_path)
                img.width, img.height = int(img.width * 0.40), int(img.height * 0.40)
                ws2.add_image(img, anchor_cells[idx])

    absolute_path = os.path.abspath(output_filename)
    wb.save(output_filename)
    print(f"\n[SUCCESSO] Report Excel salvato in: {absolute_path}\n")